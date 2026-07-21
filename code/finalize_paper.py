# -*- coding: utf-8 -*-
"""
Finalize ECERAG_Conference_Paper_REVISED.docx for submission:
  - IEEE US Letter page size + margins
  - proper two-column section layout
  - author/affiliation placeholders clearly marked (not silently-submittable dummy text)
  - a code/data availability statement pointing at the GitHub repo
"""
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches

PATH = r"C:\Users\yomis\Downloads\ieee\ECERAG_Conference_Paper_REVISED.docx"
REPO_URL = "https://github.com/yomisys/LLM"
# pinned so the cited artifact matches exactly what this paper describes (40-item
# benchmark, 0.5B generator) even though the repo's main branch may move ahead
# with a larger follow-up benchmark/experiment
REPO_COMMIT = "bd89046"


def set_ieee_letter_two_column(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)

    sectPr = section._sectPr
    existing = sectPr.find(qn("w:cols"))
    if existing is not None:
        sectPr.remove(existing)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "288")  # 0.2in gutter, in twentieths of a point (twips)
    sectPr.append(cols)


def set_paragraph_text(p, text, bold=False, italic=False):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.bold = bold
    run.italic = italic


def main():
    d = docx.Document(PATH)

    set_ieee_letter_two_column(d)

    paras = d.paragraphs
    # paragraph 1 = author line, paragraph 2 = affiliation block (see earlier structural dump)
    set_paragraph_text(paras[1], "[AUTHOR NAME(S) -- TO BE FILLED IN BEFORE SUBMISSION]", bold=True)
    set_paragraph_text(paras[2], "[DEPARTMENT / INSTITUTION -- TO BE FILLED IN]\n"
                                  "[CITY, COUNTRY -- TO BE FILLED IN]\n"
                                  "[EMAIL -- TO BE FILLED IN]")

    # availability statement, inserted right after the Conclusion section (before References)
    def find(substr):
        for i, p in enumerate(d.paragraphs):
            if substr in p.text:
                return i
        raise ValueError(substr)

    idx_conclusion = find("Future work will extend ECERAG")
    anchor = d.paragraphs[idx_conclusion + 1]
    p = anchor.insert_paragraph_before("", style=anchor.style)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(
        f"Code and Data Availability: The ECERAG pipeline implementation, corpus sources, "
        f"the 40-item benchmark, and experiment results reported in this paper are available "
        f"at {REPO_URL} (commit {REPO_COMMIT}). The repository's main branch may include a "
        f"subsequent, larger benchmark and additional experiments beyond what is reported here."
    )
    run.font.name = "Times New Roman"
    run.italic = True

    d.save(PATH)
    print(f"Finalized {PATH}")


if __name__ == "__main__":
    main()
